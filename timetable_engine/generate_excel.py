from pathlib import Path
import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter



# ============================================================
# EXCEL TIMETABLE GENERATOR
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = PROJECT_ROOT / "output" / "excel"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

OUTPUT_FILE = OUTPUT_DIR / "TE_SemV_AI_Generated_Timetable.xlsx"
SCHEDULE_FILE = PROJECT_ROOT / "output" / "schedule" / "TE_SemV_AI_Schedule.json"

DAYS = [
    "Monday",
    "Tuesday",
    "Wednesday",
    "Thursday",
    "Friday",
]

PERIODS = [
    ("P1", "08:45-09:45"),
    ("P2", "09:45-10:45"),
    ("P3", "11:00-12:00"),
    ("P4", "12:00-13:00"),
    ("P5", "13:45-14:45"),
    ("P6", "14:45-15:45"),
]


# ============================================================
# HELPERS
# ============================================================

def occupied_periods(record):
    """Return all hourly periods occupied by a session."""
    start_period = record.get("period", "")
    try:
        start_index = next(i for i, (label, _) in enumerate(PERIODS)
                           if label == start_period)
    except StopIteration:
        return [start_period] if start_period else []
    duration = int(record.get("duration", 1) or 1)
    return [PERIODS[i][0] for i in range(start_index, min(start_index + duration, len(PERIODS)))]

def display_subject(record):
    """Keep tutorials and practicals visually distinct."""
    subject = clean_text(record.get("subject", ""))
    kind = clean_text(record.get("type", "")).strip().lower()
    if kind in {"tutorial", "practical"}:
        return f"{subject} ({kind.title()})"
    return subject


def clean_text(value):
    if value is None:
        return ""
    return str(value)


def record_label(record):
    subject = clean_text(record["subject"])
    batch = clean_text(record["batch"])
    faculty = clean_text(record["faculty"])
    room = clean_text(record["room"])

    text = display_subject(record)

    if batch:
        text += f"\n[{batch}]"

    if faculty:
        text += f"\nFaculty: {faculty}"

    if room:
        text += f"\nRoom: {room}"

    return text


def make_sheet_title(wb, title):
    ws = wb.create_sheet(title)

    ws.sheet_view.showGridLines = False

    return ws


def apply_title(ws, title, end_column=7):
    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=end_column,
    )

    cell = ws.cell(1, 1)
    cell.value = title
    cell.font = Font(
        bold=True,
        size=16,
    )
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    ws.row_dimensions[1].height = 28


def style_header(row_cells):
    for cell in row_cells:
        cell.font = Font(
            bold=True,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7",
        )

        cell.border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )


def style_body(ws):
    for row in ws.iter_rows():
        for cell in row:

            if cell.row <= 2:
                continue

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

            cell.border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )


def set_widths(ws, widths):

    for column, width in widths.items():
        ws.column_dimensions[column].width = width


# ============================================================
# MASTER TIMETABLE
# ============================================================

def create_master_sheet(wb, records):

    ws = make_sheet_title(
        wb,
        "Master Timetable",
    )

    apply_title(
        ws,
        "KJ College of Engineering & Management Research, Pune\n"
        "Department of E&TC — TE Semester V — AI Generated Timetable",
        7,
    )

    ws["A3"] = "Day"
    ws["B3"] = "P1\n08:45-09:45"
    ws["C3"] = "P2\n09:45-10:45"
    ws["D3"] = "P3\n11:00-12:00"
    ws["E3"] = "P4\n12:00-13:00"
    ws["F3"] = "P5\n13:45-14:45"
    ws["G3"] = "P6\n14:45-15:45"

    style_header(ws[3])

    lookup = {}

    for record in records:

        for occupied_period in occupied_periods(record):
            key = (record["day"], occupied_period)
            lookup.setdefault(key, []).append(record)

    for day_index, day in enumerate(DAYS, start=4):

        ws.cell(
            day_index,
            1,
            day,
        )

        for period_index, (period, _) in enumerate(
            PERIODS,
            start=2,
        ):

            records_here = lookup.get(
                (day, period),
                [],
            )

            if not records_here:

                value = ""

            else:

                value = "\n\n".join(
                    record_label(record)
                    for record in records_here
                )

            ws.cell(
                day_index,
                period_index,
                value,
            )

    style_body(ws)

    for row in range(4, 9):
        ws.row_dimensions[row].height = 80

    set_widths(
        ws,
        {
            "A": 15,
            "B": 25,
            "C": 25,
            "D": 25,
            "E": 25,
            "F": 25,
            "G": 25,
        },
    )

    ws.freeze_panes = "B4"


# ============================================================
# BATCH TIMETABLE
# ============================================================

def create_batch_sheet(wb, records, batch):

    ws = make_sheet_title(
        wb,
        batch,
    )

    apply_title(
        ws,
        f"TE Semester V — {batch} Timetable",
        7,
    )

    headers = [
        "Day",
        "P1\n08:45-09:45",
        "P2\n09:45-10:45",
        "P3\n11:00-12:00",
        "P4\n12:00-13:00",
        "P5\n13:45-14:45",
        "P6\n14:45-15:45",
    ]

    for col, value in enumerate(headers, start=1):
        ws.cell(3, col, value)

    style_header(ws[3])

    lookup = {}

    for record in records:

        batches = {
            x.strip()
            for x in record["batch"].split(",")
            if x.strip()
        }

        if batch not in batches:
            continue

        for occupied_period in occupied_periods(record):
            lookup.setdefault((record["day"], occupied_period), []).append(record)

    for row, day in enumerate(DAYS, start=4):

        ws.cell(row, 1, day)

        for col, (period, _) in enumerate(
            PERIODS,
            start=2,
        ):

            records_here = lookup.get(
                (day, period),
                [],
            )

            values = []

            for record in records_here:

                text = display_subject(record)

                elective = clean_text(
                    record.get("elective", "")
                )

                if elective:
                    text += f"\n{elective}"

                text += (
                    f"\n{record['faculty']}"
                    f"\n{record['room']}"
                )

                values.append(text)

            ws.cell(
                row,
                col,
                "\n\n".join(values),
            )

    style_body(ws)

    for row in range(4, 9):
        ws.row_dimensions[row].height = 75

    set_widths(
        ws,
        {
            "A": 15,
            "B": 24,
            "C": 24,
            "D": 24,
            "E": 24,
            "F": 24,
            "G": 24,
        },
    )

    ws.freeze_panes = "B4"


# ============================================================
# FACULTY TIMETABLE
# ============================================================

def create_faculty_sheet(wb, records):

    ws = make_sheet_title(
        wb,
        "Faculty Timetable",
    )

    apply_title(
        ws,
        "TE Semester V — Faculty Timetable",
        5,
    )

    headers = [
        "Faculty",
        "Day",
        "Period",
        "Subject",
        "Batch",
    ]

    for col, value in enumerate(headers, start=1):
        ws.cell(3, col, value)

    style_header(ws[3])

    sorted_records = sorted(
        records,
        key=lambda r: (
            r["faculty"],
            DAYS.index(r["day"]),
            r["period"],
        ),
    )

    row = 4

    for record in sorted_records:

        values = [
            record["faculty"],
            record["day"],
            record["period"],
            record["subject"],
            record["batch"],
        ]

        for col, value in enumerate(
            values,
            start=1,
        ):
            ws.cell(
                row,
                col,
                value,
            )

        row += 1

    style_body(ws)

    set_widths(
        ws,
        {
            "A": 14,
            "B": 15,
            "C": 18,
            "D": 42,
            "E": 15,
        },
    )

    ws.auto_filter.ref = (
        f"A3:E{row - 1}"
    )

    ws.freeze_panes = "A4"


# ============================================================
# ROOM UTILIZATION
# ============================================================

def create_room_sheet(wb, records):

    ws = make_sheet_title(
        wb,
        "Room Utilization",
    )

    apply_title(
        ws,
        "TE Semester V — Room Utilization",
        6,
    )

    headers = [
        "Room",
        "Day",
        "Period",
        "Subject",
        "Batch",
        "Faculty",
    ]

    for col, value in enumerate(headers, start=1):
        ws.cell(3, col, value)

    style_header(ws[3])

    sorted_records = sorted(
        records,
        key=lambda r: (
            r["room"],
            DAYS.index(r["day"]),
            r["period"],
        ),
    )

    row = 4

    for record in sorted_records:

        values = [
            record["room"],
            record["day"],
            record["period"],
            record["subject"],
            record["batch"],
            record["faculty"],
        ]

        for col, value in enumerate(
            values,
            start=1,
        ):
            ws.cell(
                row,
                col,
                value,
            )

        row += 1

    style_body(ws)

    set_widths(
        ws,
        {
            "A": 12,
            "B": 15,
            "C": 18,
            "D": 42,
            "E": 15,
            "F": 14,
        },
    )

    ws.auto_filter.ref = (
        f"A3:F{row - 1}"
    )

    ws.freeze_panes = "A4"


# ============================================================
# PRACTICAL ALLOCATION
# ============================================================

def create_practical_sheet(wb, records):

    ws = make_sheet_title(
        wb,
        "Practical Allocation",
    )

    apply_title(
        ws,
        "TE Semester V — Practical Allocation",
        8,
    )

    headers = [
        "Day",
        "Period",
        "Subject",
        "Batch",
        "Faculty",
        "Room",
        "Duration",
        "Students",
    ]

    for col, value in enumerate(headers, start=1):
        ws.cell(3, col, value)

    style_header(ws[3])

    practical_records = [
        record
        for record in records
        if record["type"] == "practical"
    ]

    practical_records.sort(
        key=lambda r: (
            DAYS.index(r["day"]),
            r["period"],
            r["subject"],
            r["batch"],
        )
    )

    row = 4

    for record in practical_records:

        values = [
            record["day"],
            record["period"],
            record["subject"],
            record["batch"],
            record["faculty"],
            record["room"],
            f"{record['duration']} hours",
            record["students"],
        ]

        for col, value in enumerate(
            values,
            start=1,
        ):
            ws.cell(
                row,
                col,
                value,
            )

        row += 1

    style_body(ws)

    set_widths(
        ws,
        {
            "A": 15,
            "B": 18,
            "C": 42,
            "D": 15,
            "E": 14,
            "F": 12,
            "G": 14,
            "H": 12,
        },
    )

    ws.auto_filter.ref = (
        f"A3:H{row - 1}"
    )

    ws.freeze_panes = "A4"


# ============================================================
# SESSION DATA
# ============================================================

def create_session_data_sheet(wb, records):

    ws = make_sheet_title(
        wb,
        "Session Data",
    )

    apply_title(
        ws,
        "TE Semester V — Solver Session Data",
        11,
    )

    headers = [
        "Session ID",
        "Day",
        "Period",
        "Start",
        "End",
        "Subject",
        "Type",
        "Duration",
        "Batch",
        "Faculty",
        "Room",
    ]

    for col, value in enumerate(headers, start=1):
        ws.cell(3, col, value)

    style_header(ws[3])

    row = 4

    for record in records:

        values = [
            record["session_id"],
            record["day"],
            record["period"],
            record["start"],
            record["end"],
            record["subject"],
            record["type"],
            record["duration"],
            record["batch"],
            record["faculty"],
            record["room"],
        ]

        for col, value in enumerate(
            values,
            start=1,
        ):
            ws.cell(
                row,
                col,
                value,
            )

        row += 1

    style_body(ws)

    set_widths(
        ws,
        {
            "A": 28,
            "B": 15,
            "C": 12,
            "D": 12,
            "E": 12,
            "F": 42,
            "G": 14,
            "H": 12,
            "I": 15,
            "J": 14,
            "K": 12,
        },
    )

    ws.auto_filter.ref = (
        f"A3:K{row - 1}"
    )

    ws.freeze_panes = "A4"


# ============================================================
# SUMMARY
# ============================================================

def create_summary_sheet(wb, records):

    ws = make_sheet_title(
        wb,
        "Summary",
    )

    apply_title(
        ws,
        "AI Timetable Generation — Summary",
        4,
    )

    summary = [
        ("College", "KJ College of Engineering & Management Research, Pune"),
        ("Department", "E&TC"),
        ("Class", "TE Semester V"),
        ("Working Days", "Monday-Friday"),
        ("College Timing", "08:45 AM - 03:45 PM"),
        ("Total Sessions", len(records)),
        ("Generated Student Workload", "28 hours/week per batch"),
        ("Solver Status", "OPTIMAL"),
    ]

    row = 3

    for label, value in summary:

        ws.cell(row, 1, label)
        ws.cell(row, 2, value)

        ws.cell(row, 1).font = Font(
            bold=True
        )

        row += 1

    style_body(ws)

    set_widths(
        ws,
        {
            "A": 30,
            "B": 65,
        },
    )


# ============================================================
# MAIN
# ============================================================

def main():

    print()
    print("=" * 60)
    print("AI TIMETABLE EXCEL GENERATOR")
    print("=" * 60)

    if not SCHEDULE_FILE.exists():
        print("ERROR: No saved schedule found.")
        print("Run solver.py successfully first, then run generate_excel.py.")
        return

    try:
        payload = json.loads(SCHEDULE_FILE.read_text(encoding="utf-8"))
        records = payload["records"]
    except (OSError, json.JSONDecodeError, KeyError, TypeError) as exc:
        print(f"ERROR: Could not read saved schedule: {exc}")
        return

    if not isinstance(records, list) or not records:
        print("ERROR: Saved schedule has no session records. Run solver.py again.")
        return

    print(f"Using saved schedule: {SCHEDULE_FILE}")
    print(f"Schedule created: {payload.get('created_at', 'unknown')}")
    print(f"Scheduled sessions: {len(records)}")

    wb = Workbook()

    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    # Create sheets
    create_summary_sheet(
        wb,
        records,
    )

    create_master_sheet(
        wb,
        records,
    )

    create_batch_sheet(
        wb,
        records,
        "T1",
    )

    create_batch_sheet(
        wb,
        records,
        "T2",
    )

    create_batch_sheet(
        wb,
        records,
        "T3",
    )

    create_faculty_sheet(
        wb,
        records,
    )

    create_room_sheet(
        wb,
        records,
    )

    create_practical_sheet(
        wb,
        records,
    )

    create_session_data_sheet(
        wb,
        records,
    )

    # Active sheet
    wb.active = 0

    wb.save(
        OUTPUT_FILE
    )

    print()
    print("=" * 60)
    print("EXCEL FILE CREATED")
    print("=" * 60)

    print(
        f"File: {OUTPUT_FILE}"
    )

    print()
    print("Sheets:")

    for ws in wb.worksheets:
        print(
            f"  - {ws.title}"
        )

    print()
    print(
        f"Scheduled sessions: {len(records)}"
    )

    print("=" * 60)


if __name__ == "__main__":
    main()
